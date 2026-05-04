import openpyxl
import os
import warnings
from datetime import datetime
import calendar
from openpyxl.styles import Font
import sys

# 1. SEGURIDAD Y CONFIGURACIÓN DE ALERTAS
warnings.filterwarnings("ignore", category=UserWarning)

# =================================================================
# BLOQUE DE CONFIGURACIÓN DINÁMICA
# =================================================================
# Este script automatiza el procesamiento de reportes financieros.
# Detecta automáticamente los archivos según palabras clave.
# =================================================================

CARPETA_BASE = os.path.dirname(os.path.abspath(sys.argv[0]))

# Etiquetas genéricas para búsqueda de archivos
PALABRA_CLAVE_DIFERIDO = "Diferido"
PALABRA_CLAVE_REPORTE  = "reporte"

RUTA_ENTRADA = CARPETA_BASE
RUTA_SALIDA  = os.path.join(CARPETA_BASE, "Resultado_Procesado.xlsx")

# =================================================================
# FUNCIONES DE UTILIDAD
# =================================================================

def extraer_mes_de_nombre(nombre_archivo):
    """Identifica el mes de trabajo basado en el nombre del archivo."""
    meses = {
        "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, 
        "mayo": 5, "junio": 6, "julio": 7, "agosto": 8, 
        "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12
    }
    nombre_min = nombre_archivo.lower()
    for mes_texto, mes_num in meses.items():
        if mes_texto in nombre_min:
            return mes_num
    return None

def buscar_archivo_mas_reciente(carpeta, palabra_clave):
    """Localiza el archivo .xlsx más nuevo que contenga la palabra clave."""
    archivos = [f for f in os.listdir(carpeta) if palabra_clave.lower() in f.lower() and f.endswith('.xlsx') and not f.startswith('~$')]
    if not archivos:
        return None
    # Ordenar por fecha de modificación (el más reciente primero)
    archivos.sort(key=lambda x: os.path.getmtime(os.path.join(carpeta, x)), reverse=True)
    return archivos[0]

def obtener_ultimo_dia_mes(anio, mes):
    """Calcula el último día calendario de un mes específico."""
    ultimo_dia = calendar.monthrange(anio, mes)[1]
    return datetime(anio, mes, ultimo_dia)

# =================================================================
# PROCESO PRINCIPAL DE AUTOMATIZACIÓN
# =================================================================

def ejecutar_motor_procesamiento():
    # --- 1. BÚSQUEDA AUTOMÁTICA DE ARCHIVOS ---
    archivo_dif = buscar_archivo_mas_reciente(RUTA_ENTRADA, PALABRA_CLAVE_DIFERIDO)
    archivo_rep = buscar_archivo_mas_reciente(RUTA_ENTRADA, PALABRA_CLAVE_REPORTE)

    if not archivo_dif or not archivo_rep:
        print(f"\n[ERROR]: No se encontraron archivos necesarios en: {RUTA_ENTRADA}")
        return

    # Detección de periodo de trabajo
    MES_DETECTADO = extraer_mes_de_nombre(archivo_rep)
    ANIO_DETECTADO = datetime.now().year

    if not MES_DETECTADO:
        print(f"[ERROR]: No se pudo identificar el mes en el archivo: {archivo_rep}")
        return

    # --- 2. CONFIRMACIÓN DE PARÁMETROS ---
    print(f"\n" + "="*40)
    print(f"ARCHIVO ORIGEN 1: {archivo_dif}")
    print(f"ARCHIVO ORIGEN 2: {archivo_rep}")
    print(f"PERIODO DETECTADO: {MES_DETECTADO}/{ANIO_DETECTADO}")
    print("="*40)

    confirmacion = input("¿Desea iniciar el proceso con estos datos? (s/n): ").strip().lower()
    if confirmacion != 's':
        print("Operación cancelada.")
        return

    path_diferido = os.path.join(RUTA_ENTRADA, archivo_dif)
    path_reporte = os.path.join(RUTA_ENTRADA, archivo_rep)

    # Configuración de fechas de corte
    fecha_limite = obtener_ultimo_dia_mes(ANIO_DETECTADO, MES_DETECTADO)
    mes_sig = (MES_DETECTADO % 12) + 1
    anio_sig = ANIO_DETECTADO if mes_sig > 1 else ANIO_DETECTADO + 1
    fecha_corte_mes_sig = datetime(anio_sig, mes_sig, 1)

    print(f"-> Iniciando análisis para el cierre de {MES_DETECTADO}/{ANIO_DETECTADO}...")

    try:
        wb = openpyxl.load_workbook(path_diferido)

        # --- 3. DETECCIÓN DE HOJA DE HISTÓRICO ---
        nombre_hoja_hist = None
        for nombre in wb.sheetnames:
            if "-" in nombre and nombre.strip().startswith("20"):
                nombre_hoja_hist = nombre
                break

        if not nombre_hoja_hist:
            print("[ERROR]: No se encontró la hoja de histórico anual.")
            return

        # --- 4. LIMPIEZA Y PREPARACIÓN DE ESTRUCTURA ---
        for nombre in wb.sheetnames:
            ws_t = wb[nombre]
            if hasattr(ws_t, '_images'): ws_t._images = []
            # Eliminación de hojas temporales previas
            if any(nombre.startswith(p) for p in ["NC", "Ftbaj", "Res.Dev"]):
                wb.remove(ws_t)

        ws_reporte = wb["Reporte"] if "Reporte" in wb.sheetnames else wb.create_sheet("Reporte")
        ws_reporte.delete_rows(1, ws_reporte.max_row + 100)

        ws_dev = wb["Dev. no fact."]
        if ws_dev.max_row > 1: ws_dev.delete_rows(2, ws_dev.max_row + 100)

        ws_fact = wb["fact. no dev."]
        if ws_fact.max_row > 1: ws_fact.delete_rows(2, ws_fact.max_row + 100)

        # --- 5. CARGA E INTEGRACIÓN DE DATOS ---
        wb_orig = openpyxl.load_workbook(path_reporte, data_only=True)
        ws_orig = wb_orig[wb_orig.sheetnames[0]]

        for row in ws_orig.iter_rows(values_only=True):
            if any(cell is not None for cell in row): 
                ws_reporte.append(row)

        filas_reporte = list(ws_reporte.iter_rows(min_row=2, values_only=True))
        filas_reporte.sort(key=lambda x: x[2] if isinstance(x[2], datetime) else datetime.min, reverse=True)

        ws_reporte.delete_rows(2, ws_reporte.max_row)
        for r in filas_reporte: ws_reporte.append(r)

        # --- 6. DISTRIBUCIÓN DINÁMICA POR CRITERIOS CONTABLES ---
        data_para_historico = []
        for r_data in filas_reporte:
            l_row = list(r_data)
            while len(l_row) < 26: l_row.append(None)

            currency = str(l_row[14]).strip().upper() 
            monto_p = l_row[15] or 0                  

            l_row[24] = monto_p if "PEN" in currency else None
            l_row[25] = monto_p if "USD" in currency else None

            # Filtrado por lógica de periodos
            if isinstance(r_data[2], datetime) and r_data[2].month == mes_sig:
                if not (isinstance(r_data[4], datetime) and r_data[4].month == mes_sig):
                    ws_dev.append(l_row)
            elif isinstance(r_data[2], datetime) and r_data[2].month == MES_DETECTADO:
                if isinstance(r_data[5], datetime) and r_data[5] >= fecha_corte_mes_sig:
                    ws_fact.append(l_row)
                    data_para_historico.append(tuple(l_row))

        # --- 7. CÁLCULO DE TOTALES ---
        ultima_fila_dev = ws_dev.max_row + 1
        suma_dev_pen = 0
        suma_dev_usd = 0

        if ultima_fila_dev > 2:
            ws_dev.cell(row=ultima_fila_dev, column=24, value="TOTALES:").font = Font(bold=True)
            suma_dev_pen = sum(row[24].value for row in ws_dev.iter_rows(min_row=2, max_row=ultima_fila_dev-1) if isinstance(row[24].value, (int, float)))
            suma_dev_usd = sum(row[25].value for row in ws_dev.iter_rows(min_row=2, max_row=ultima_fila_dev-1) if isinstance(row[25].value, (int, float)))
            ws_dev.cell(row=ultima_fila_dev, column=25, value=suma_dev_pen).font = Font(bold=True)
            ws_dev.cell(row=ultima_fila_dev, column=26, value=suma_dev_usd).font = Font(bold=True)

        # --- 8. ACTUALIZACIÓN DE HOJA HISTÓRICA ---
        if nombre_hoja_hist in wb.sheetnames:
            ws_hist = wb[nombre_hoja_hist] 
            
            # Limpieza de columnas de cálculo
            for row in ws_hist.iter_rows(min_row=2, min_col=25, max_col=26):
                for cell in row: cell.value = None

            filas_hist = list(ws_hist.iter_rows(min_row=2, values_only=True))
            ws_hist.delete_rows(2, ws_hist.max_row + 100)

            # Evitar duplicados del periodo actual
            for r_h in filas_hist:
                if r_h[5] and isinstance(r_h[5], datetime) and r_h[5].month == MES_DETECTADO and r_h[5].year == ANIO_DETECTADO:
                    continue
                ws_hist.append(r_h)
            for r_new in data_para_historico: ws_hist.append(r_new)

            config_intervalo = {"mensual": 30, "trimestral": 90, "semestral": 180, "anual": 365, "bianual": 730}

            for i in range(2, ws_hist.max_row + 1):
                f_fin = ws_hist.cell(row=i, column=6).value
                intervalo = str(ws_hist.cell(row=i, column=18).value).strip().lower()

                ws_hist.cell(row=i, column=30, value=fecha_limite) # Fecha de Corte

                curr = str(ws_hist.cell(row=i, column=15).value).strip().upper()
                monto = ws_hist.cell(row=i, column=16).value
                if curr == "PEN": ws_hist.cell(row=i, column=25, value=monto)
                elif curr == "USD": ws_hist.cell(row=i, column=26, value=monto)

                # Lógica de Diferimiento
                resultado_days = 0
                if isinstance(f_fin, datetime):
                    dias_base = max(0, (f_fin - fecha_limite).days)
                    tope = config_intervalo.get(intervalo, 30)
                    resultado_days = min(dias_base, tope)
                    ws_hist.cell(row=i, column=29, value=resultado_days)
                
                divisor = config_intervalo.get(intervalo)
                if divisor and resultado_days > 0:
                    m_p = ws_hist.cell(row=i, column=25).value
                    m_u = ws_hist.cell(row=i, column=26).value
                    if m_p: ws_hist.cell(row=i, column=27, value=(m_p / divisor) * resultado_days)
                    if m_u: ws_hist.cell(row=i, column=28, value=(m_u / divisor) * resultado_days)

            # Resumen final de la hoja histórica
            ultima_fila = ws_hist.max_row + 1
            ws_hist.cell(row=ultima_fila, column=26, value="TOTALES:").font = Font(bold=True)
            suma_pen = sum(ws_hist.cell(row=r, column=27).value for r in range(2, ultima_fila) if isinstance(ws_hist.cell(row=r, column=27).value, (int, float)))
            suma_usd = sum(ws_hist.cell(row=r, column=28).value for r in range(2, ultima_fila) if isinstance(ws_hist.cell(row=r, column=28).value, (int, float)))
            ws_hist.cell(row=ultima_fila, column=27, value=suma_pen).font = Font(bold=True)
            ws_hist.cell(row=ultima_fila, column=28, value=suma_usd).font = Font(bold=True)

        # --- 9. GENERACIÓN DE REPORTES EJECUTIVOS ---
        ws_final = wb["RESUMEN"] if "RESUMEN" in wb.sheetnames else wb.create_sheet("RESUMEN")
        ws_final.delete_rows(1, ws_final.max_row + 100)

        meses_nombres = {1:"Enero", 2:"Febrero", 3:"Marzo", 4:"Abril", 5:"Mayo", 6:"Junio",
                         7:"Julio", 8:"Agosto", 9:"Septiembre", 10:"Octubre", 11:"Noviembre", 12:"Diciembre"}
        
        ws_final.cell(row=1, column=1, value=f"Cierre {meses_nombres.get(MES_DETECTADO)} {ANIO_DETECTADO}").font = Font(bold=True, size=14)
        
        # Estructura del Resumen Financiero
        items = [
            ("Facturado pendiente de devengo", suma_pen, suma_usd),
            ("Devengado pendiente de factura", suma_dev_pen, suma_dev_usd)
        ]

        for idx, (label, p, u) in enumerate(items, start=5):
            ws_final.cell(row=idx, column=1, value=label)
            ws_final.cell(row=idx, column=2, value=p).number_format = '"S/ " #,##0.00'
            ws_final.cell(row=idx, column=3, value=u).number_format = '"$ " #,##0.00'

        wb.save(RUTA_SALIDA)
        print(f"\n[ÉXITO]: Reporte generado en: {RUTA_SALIDA}")

    except Exception as e:
        print(f"\n[ERROR CRÍTICO]: {str(e)}")

if __name__ == "__main__":
    ejecutar_motor_procesamiento()
    input("\nPresiona ENTER para finalizar...")